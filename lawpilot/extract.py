"""Extract text from uploaded files, keeping page numbers where the format has them."""
from __future__ import annotations

import io
import re
from typing import List, Optional, Tuple

Pages = List[Tuple[Optional[int], str]]

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}


class ExtractionError(Exception):
    pass


def extract(data: bytes, filename: str) -> Pages:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in SUPPORTED:
        raise ExtractionError(f"Unsupported file type '{ext or filename}'. Use PDF, DOCX, TXT, MD or HTML.")
    if ext == ".pdf":
        pages = _pdf(data)
    elif ext == ".docx":
        pages = _docx(data)
    elif ext in (".html", ".htm"):
        pages = [(None, html_to_text(data.decode("utf-8", errors="replace")))]
    else:
        pages = [(None, data.decode("utf-8", errors="replace"))]
    pages = [(p, clean(t)) for p, t in pages if t and t.strip()]
    if not pages or sum(len(t) for _, t in pages) < 40:
        if ext == ".pdf":
            raise ExtractionError("No text could be extracted. This looks like a scanned PDF; run OCR on it first "
                                  "(for example with ocrmypdf) and upload the result.")
        raise ExtractionError("The file contains no readable text.")
    if ext == ".pdf" and not readable(pages):
        raise ExtractionError("The PDF's text layer is unreadable (mostly symbols). This usually means a scanned "
                              "document or a Hindi edition with an unsupported font. Use the English text PDF, "
                              "or run OCR on it first.")
    return pages


def readable(pages: Pages) -> bool:
    """False when 'text' is mostly punctuation, e.g. a Devanagari font with no Unicode mapping."""
    text = "".join(t for _, t in pages)
    visible = [ch for ch in text if not ch.isspace()]
    if len(visible) < 2000:
        return True
    return sum(ch.isalpha() for ch in visible) / len(visible) >= 0.5


def _pdf(data: bytes) -> Pages:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("PDF support requires the 'pypdf' package.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError("This PDF is password protected.") from exc
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not read PDF: {exc}") from exc


def _docx(data: bytes) -> Pages:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError("DOCX support requires the 'python-docx' package.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [(None, "\n\n".join(parts))]


def html_to_text(html: str) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return re.sub(r"<[^>]+>", " ", html)
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "noscript"]):
        tag.decompose()
    main = soup.find("div", class_="judgments") or soup.find("main") or soup.body or soup
    return main.get_text("\n")


def clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
